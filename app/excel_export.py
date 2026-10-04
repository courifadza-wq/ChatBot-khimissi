# ============================================================
#  Génération d'un fichier Excel (.xlsx) pour une commande
#  Chaque article de la commande devient une ligne du tableau.
# ============================================================
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _style_header(ws, ncols):
    """Applique un style aux en-têtes (colonnes 1..ncols)."""
    fill = PatternFill("solid", fgColor="1E3A5F")
    font = Font(bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border


def _autofit(ws, nrows, ncols):
    """Ajuste la largeur des colonnes à la longueur du contenu."""
    for col in range(1, ncols + 1):
        max_len = 8
        for row in range(1, nrows + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 2, 45)


def order_to_excel_bytes(order) -> bytes:
    """
    Construit un fichier .xlsx (en mémoire) décrivant la commande.
    Renvoie les octets du fichier prêt à être joint à l'email.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Commande"

    headers = ["Produit", "Quantité", "Prix unitaire (DZD)", "Total (DZD)"]
    ws.append(headers)
    _style_header(ws, len(headers))

    for it in order.items:
        price = it.get("price", 0)
        ws.append([it["name"], it["qty"], price, price * it["qty"]])

    # Ligne total
    ws.append([])
    ws.append(["TOTAL", "", "", order.total])
    total_row = ws.max_row
    ws.cell(row=total_row, column=4).font = Font(bold=True)

    # Infos client en dessous
    ws.append([])
    ws.append(["Client", order.name])
    ws.append(["Téléphone", order.phone])
    ws.append(["Adresse", order.address])
    ws.append(["Paiement", order.payment])
    ws.append(["Date", datetime.now().strftime("%Y-%m-%d %H:%M")])

    _autofit(ws, ws.max_row, len(headers))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
